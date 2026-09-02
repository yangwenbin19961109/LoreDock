"""Generate deterministic, redistributable PDF and DOCX retrieval fixtures."""

from __future__ import annotations

import os
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "evals" / "fixtures" / "documents"


def _set_word_font(run: object, name: str, size: float, *, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def build_docx(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 14, 7),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    _set_word_font(title.add_run("LoreDock 本地模型运行手册"), "Microsoft YaHei", 22, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    _set_word_font(subtitle.add_run("Phase 2.5 DOCX 检索评测资料"), "Microsoft YaHei", 11)

    sections = [
        ("模型状态", "模型状态分为 missing、downloading、verifying、ready 和 failed。只有 ready 状态允许创建 ONNX 推理会话。"),
        ("下载临时文件", "下载内容先写入扩展名为 .part 的临时文件。网络中断时临时文件不能被模型扫描器识别为正式资产。"),
        ("完整性校验", "模型文件必须同时验证预期字节数和 SHA-256。任一条件不匹配都进入 failed 状态。"),
        ("不可变版本", "模型清单固定 repository revision，不能在相同模型标识下静默追踪 main 分支的新文件。"),
        ("查询前缀", "multilingual-e5-small 的查询文本必须添加 query: 前缀，资料文本必须添加 passage: 前缀。"),
        ("池化规则", "ONNX 输出按照 attention mask 执行 mean pooling，随后对每个 384 维向量做 L2 归一化。"),
        ("长度限制", "E5 输入最多 512 tokens。超长资料必须在分块阶段控制，模型层仍执行确定性截断。"),
        ("批处理", "默认推理批大小是 16。批大小可以根据设备内存调整，但不能改变向量语义或索引契约。"),
        ("CPU Provider", "轻量桌面默认使用 ONNX Runtime CPUExecutionProvider，不要求用户安装 CUDA、PyTorch 或系统模型服务。"),
        ("模型切换", "模型标识、校验和、维度、前缀、池化或归一化规则变化时，所有相关文档向量都必须重建。"),
        ("失败降级", "模型加载或推理失败时搜索回退到 BM25-only。已有 FTS 索引继续可用，并向用户显示可恢复错误。"),
        ("本地隐私", "本地模型推理不发送文档正文和查询到外部提供商。日志也不记录完整输入文本。"),
    ]
    for heading, body in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
    document.core_properties.title = "LoreDock 本地模型运行手册"
    document.core_properties.subject = "Non-sensitive retrieval evaluation fixture"
    document.core_properties.author = "LoreDock"
    document.core_properties.created = datetime(2026, 9, 2, tzinfo=UTC)
    document.core_properties.modified = datetime(2026, 9, 2, tzinfo=UTC)
    document.save(path)

    normalized = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(2026, 9, 2, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    os.replace(normalized, path)


def build_pdf(path: Path) -> None:
    noto_path = Path(r"C:\Windows\Fonts\NotoSansSC-VF.ttf")
    if not noto_path.exists():
        raise RuntimeError("Noto Sans SC is required to regenerate the PDF fixture")
    pdfmetrics.registerFont(TTFont("NotoSansSC", str(noto_path)))
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "FixtureTitle",
        parent=styles["Title"],
        fontName="NotoSansSC",
        fontSize=22,
        leading=28,
        textColor=HexColor("#1F4D78"),
        alignment=TA_CENTER,
        spaceAfter=16,
    )
    heading = ParagraphStyle(
        "FixtureHeading",
        parent=styles["Heading1"],
        fontName="NotoSansSC",
        fontSize=14,
        leading=19,
        textColor=HexColor("#2E74B5"),
        spaceBefore=12,
        spaceAfter=6,
    )
    body = ParagraphStyle(
        "FixtureBody",
        parent=styles["BodyText"],
        fontName="NotoSansSC",
        fontSize=11,
        leading=17,
        spaceAfter=8,
    )
    story = [
        Paragraph("LoreDock 故障诊断与恢复指南", title),
        Paragraph("Phase 2.5 PDF 检索评测资料", body),
        Spacer(1, 0.15 * inch),
    ]
    sections = [
        ("诊断包", "诊断包包含版本、任务状态、数据库完整性结果和脱敏日志，但不包含原始文档、认证令牌或完整用户查询。"),
        ("Core 无法启动", "先检查数据目录权限、端口占用和 app.sqlite 完整性。不要通过删除整个数据目录来尝试修复启动问题。"),
        ("模型无法加载", "依次检查 manifest、文件大小、SHA-256、ONNX Runtime 版本和 CPU Provider。模型失败不应阻塞 BM25-only。"),
        ("索引契约不匹配", "发现模型维度或 schema version 不匹配时，在旁路文件中重建派生索引，成功校验后再原子切换。"),
        ("任务长期运行", "如果 heartbeat 停止且租约过期，恢复流程把任务标记为失败或重新排队，不能让任务永久保持 running。"),
        ("数据库检查", "恢复前后执行 PRAGMA integrity_check。返回值不是 ok 时停止写入，并提示用户从一致性备份恢复。"),
        ("磁盘空间不足", "索引构建前预估模型、临时索引和 WAL 空间。磁盘不足时保留旧索引和原始资料，并删除未完成的旁路文件。"),
        ("备份恢复", "app.sqlite 使用 SQLite backup API。每库派生索引可以重建，但恢复后仍需验证 manifest 与当前模型契约一致。"),
        ("来源删除失败", "删除流程必须可重试，并记录尚未清理的精确 source ID。禁止使用宽泛通配符补偿删除。"),
        ("PDF 解析为空", "扫描型 PDF 没有文本层时标记需要 OCR。OCR 是可选组件，缺失时不影响 TXT、Markdown 和 DOCX。"),
        ("远程连接异常", "远程 MCP 失败时检查 HTTPS 证书、授权范围和 Core 健康状态，不能降级为未认证的明文连接。"),
        ("日志隐私", "日志默认记录错误码、耗时和资源 ID。文档正文、密钥、Cookie 和未脱敏查询不得进入诊断日志。"),
        ("恢复完成标准", "只有数据库完整、来源可读、索引契约匹配、搜索引用有效四项都通过，恢复任务才能标记 succeeded。"),
    ]
    for index, (section_title, section_body) in enumerate(sections):
        if index == 7:
            story.append(PageBreak())
        story.append(Paragraph(section_title, heading))
        story.append(Paragraph(section_body, body))
    document = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="LoreDock 故障诊断与恢复指南",
        author="LoreDock",
        subject="Non-sensitive retrieval evaluation fixture",
        invariant=1,
    )
    document.build(story)


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    build_docx(OUTPUT / "model-handbook.docx")
    build_pdf(OUTPUT / "recovery-guide.pdf")


if __name__ == "__main__":
    main()
