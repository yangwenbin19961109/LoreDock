"""Small, deterministic parsers used by the retrieval experiments."""

import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from loredock.ingestion.ooxml import OoxmlPackage, OoxmlPackageError, XmlElement

_HTML_BLOCK_TAGS = {
    "address",
    "article",
    "aside",
    "blockquote",
    "br",
    "dd",
    "div",
    "dl",
    "dt",
    "figcaption",
    "figure",
    "footer",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "header",
    "hr",
    "li",
    "main",
    "nav",
    "ol",
    "p",
    "pre",
    "section",
    "table",
    "tbody",
    "td",
    "tfoot",
    "th",
    "thead",
    "tr",
    "ul",
}
_HTML_IGNORED_TAGS = {"iframe", "noscript", "object", "script", "style", "template"}
_HTML_HORIZONTAL_SPACE = re.compile(r"[\t\f\v ]+")
_HTML_EXCESS_BREAKS = re.compile(r"\n[\t ]*\n(?:[\t ]*\n)+")


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    source_id: str
    title: str
    text: str
    page_spans: tuple[tuple[int, int, int], ...] = ()

    def page_at(self, char_offset: int) -> int | None:
        for page, start, end in self.page_spans:
            if start <= char_offset < end:
                return page
        return None


class _SafeHTMLTextParser(HTMLParser):
    """Extract visible HTML text without executing or fetching embedded content."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._title_parts: list[str] = []
        self._ignored_depth = 0
        self._in_title = False

    @property
    def text(self) -> str:
        normalized = "".join(self._parts).replace("\r\n", "\n").replace("\r", "\n")
        normalized = _HTML_HORIZONTAL_SPACE.sub(" ", normalized)
        normalized = re.sub(r" *\n *", "\n", normalized)
        return _HTML_EXCESS_BREAKS.sub("\n\n", normalized).strip()

    @property
    def title(self) -> str:
        return _HTML_HORIZONTAL_SPACE.sub(" ", "".join(self._title_parts)).strip()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        normalized = tag.lower()
        if normalized in _HTML_IGNORED_TAGS:
            self._ignored_depth += 1
            return
        if self._ignored_depth:
            return
        if normalized == "title":
            self._in_title = True
        elif normalized in _HTML_BLOCK_TAGS:
            self._append_break()

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        normalized = tag.lower()
        if normalized in _HTML_IGNORED_TAGS:
            if self._ignored_depth:
                self._ignored_depth -= 1
            return
        if self._ignored_depth:
            return
        if normalized == "title":
            self._in_title = False
        elif normalized in _HTML_BLOCK_TAGS:
            self._append_break()

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        if self._in_title:
            self._title_parts.append(data)
            return
        self._parts.append(data)

    def _append_break(self) -> None:
        if self._parts and not self._parts[-1].endswith("\n"):
            self._parts.append("\n")


def _parse_pdf(path: Path) -> ParsedDocument:
    parts: list[str] = []
    spans: list[tuple[int, int, int]] = []
    cursor = 0
    for page_number, page in enumerate(PdfReader(path).pages, start=1):
        text = page.extract_text() or ""
        if parts:
            parts.append("\n\n")
            cursor += 2
        start = cursor
        parts.append(text)
        cursor += len(text)
        spans.append((page_number, start, cursor))
    return ParsedDocument(path.stem, path.stem, "".join(parts), tuple(spans))


def _parse_docx(path: Path) -> ParsedDocument:
    paragraphs = [paragraph.text for paragraph in Document(str(path)).paragraphs]
    return ParsedDocument(path.stem, path.stem, "\n\n".join(paragraphs))


def _parse_html(path: Path) -> ParsedDocument:
    parser = _SafeHTMLTextParser()
    parser.feed(path.read_text(encoding="utf-8-sig"))
    parser.close()
    return ParsedDocument(path.stem, parser.title or path.stem, parser.text)


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _descendants(element: XmlElement, name: str) -> list[XmlElement]:
    matches: list[XmlElement] = []
    for child in element:
        if _local_name(child.tag) == name:
            matches.append(child)
        matches.extend(_descendants(child, name))
    return matches


def _paragraph_text(paragraph: XmlElement) -> str:
    return "".join(node.text or "" for node in _descendants(paragraph, "t")).strip()


def _table_text(table: XmlElement) -> str:
    rows: list[str] = []
    for row in _descendants(table, "tr"):
        cells: list[str] = []
        for cell in (child for child in row if _local_name(child.tag) == "tc"):
            paragraphs = [
                text
                for paragraph in _descendants(cell, "p")
                if (text := _paragraph_text(paragraph))
            ]
            cells.append(" / ".join(paragraphs).replace("|", "\\|"))
        if cells:
            rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _slide_text(slide: XmlElement) -> str:
    shape_trees = _descendants(slide, "spTree")
    if not shape_trees:
        return ""
    blocks: list[str] = []
    for shape in shape_trees[0]:
        tables = _descendants(shape, "tbl")
        if tables:
            blocks.extend(text for table in tables if (text := _table_text(table)))
            continue
        paragraphs = [
            text for paragraph in _descendants(shape, "p") if (text := _paragraph_text(paragraph))
        ]
        if paragraphs:
            blocks.append("\n".join(paragraphs))
    return "\n\n".join(blocks)


def _relationship_id(element: XmlElement) -> str | None:
    for key, value in element.attrib.items():
        if key.startswith("{") and _local_name(key) == "id":
            return value
    return None


def _parse_pptx(path: Path) -> ParsedDocument:
    presentation_part = "ppt/presentation.xml"
    with OoxmlPackage(path) as package:
        presentation = package.read_xml(presentation_part)
        relationships = {
            relationship.id: relationship
            for relationship in package.relationships(presentation_part)
            if relationship.type.rstrip("/").endswith("/slide")
        }
        slide_ids = _descendants(presentation, "sldId")
        if not slide_ids:
            raise OoxmlPackageError("PPTX presentation does not contain any slides.")

        parts: list[str] = []
        spans: list[tuple[int, int, int]] = []
        cursor = 0
        for page_number, slide_id in enumerate(slide_ids, start=1):
            relationship_id = _relationship_id(slide_id)
            relationship = relationships.get(relationship_id or "")
            if relationship is None or relationship.external:
                raise OoxmlPackageError("PPTX slide relationship is missing or external.")
            slide = package.read_xml(relationship.target)
            hidden = (slide.get("show") or "1") in {"0", "false", "False"}
            heading = f"# 幻灯片 {page_number}" + (" (隐藏)" if hidden else "")
            body = _slide_text(slide)
            text = heading if not body else f"{heading}\n\n{body}"
            if parts:
                parts.append("\n\n")
                cursor += 2
            start = cursor
            parts.append(text)
            cursor += len(text)
            spans.append((page_number, start, cursor))
    return ParsedDocument(path.stem, path.stem, "".join(parts), tuple(spans))


def _first_descendant(element: XmlElement, name: str) -> XmlElement | None:
    matches = _descendants(element, name)
    return matches[0] if matches else None


def _shared_strings(package: OoxmlPackage) -> tuple[str, ...]:
    part = "xl/sharedStrings.xml"
    if not package.has_part(part):
        return ()
    root = package.read_xml(part)
    return tuple("".join(node.text or "" for node in _descendants(item, "t")) for item in root)


def _cell_text(cell: XmlElement, shared_strings: tuple[str, ...]) -> str | None:
    cell_type = (cell.get("t") or "n").strip()
    if cell_type == "inlineStr":
        inline = _first_descendant(cell, "is")
        return (
            "" if inline is None else "".join(node.text or "" for node in _descendants(inline, "t"))
        )

    value_node = _first_descendant(cell, "v")
    raw_value = value_node.text if value_node is not None else None
    if cell_type == "s" and raw_value is not None:
        try:
            index = int(raw_value)
            return shared_strings[index]
        except (ValueError, IndexError) as error:
            raise OoxmlPackageError("XLSX cell has an invalid shared-string index.") from error
    if cell_type == "b" and raw_value is not None:
        return "TRUE" if raw_value == "1" else "FALSE"
    if raw_value is not None:
        return raw_value

    formula = _first_descendant(cell, "f")
    if formula is not None and formula.text:
        return f"={formula.text}"
    return None


def _worksheet_text(worksheet: XmlElement, shared_strings: tuple[str, ...]) -> str:
    sheet_data = _first_descendant(worksheet, "sheetData")
    if sheet_data is None:
        return ""
    values: list[str] = []
    for cell in _descendants(sheet_data, "c"):
        coordinate = (cell.get("r") or "").strip()
        value = _cell_text(cell, shared_strings)
        if value is None or value == "":
            continue
        if not coordinate:
            raise OoxmlPackageError("XLSX non-empty cell is missing its coordinate.")
        normalized = value.replace("\r\n", "\n").replace("\r", "\n")
        values.append(f"{coordinate}: {normalized}")
    return "\n".join(values)


def _parse_xlsx(path: Path) -> ParsedDocument:
    workbook_part = "xl/workbook.xml"
    with OoxmlPackage(path) as package:
        workbook = package.read_xml(workbook_part)
        relationships = {
            relationship.id: relationship
            for relationship in package.relationships(workbook_part)
            if relationship.type.rstrip("/").endswith("/worksheet")
        }
        strings = _shared_strings(package)
        sheets = _descendants(workbook, "sheet")
        if not sheets:
            raise OoxmlPackageError("XLSX workbook does not contain any worksheets.")

        parts: list[str] = []
        spans: list[tuple[int, int, int]] = []
        cursor = 0
        for page_number, sheet in enumerate(sheets, start=1):
            relationship_id = _relationship_id(sheet)
            relationship = relationships.get(relationship_id or "")
            if relationship is None or relationship.external:
                raise OoxmlPackageError("XLSX worksheet relationship is missing or external.")
            worksheet = package.read_xml(relationship.target)
            name = (sheet.get("name") or f"Sheet {page_number}").replace("\n", " ").strip()
            state = (sheet.get("state") or "visible").casefold()
            state_label = " (隐藏)" if state in {"hidden", "veryhidden"} else ""
            heading = f"# 工作表 {page_number}: {name}{state_label}"
            body = _worksheet_text(worksheet, strings)
            text = heading if not body else f"{heading}\n\n{body}"
            if parts:
                parts.append("\n\n")
                cursor += 2
            start = cursor
            parts.append(text)
            cursor += len(text)
            spans.append((page_number, start, cursor))
    return ParsedDocument(path.stem, path.stem, "".join(parts), tuple(spans))


def parse_path(
    path: Path, *, source_id: str | None = None, allowed_root: Path | None = None
) -> ParsedDocument:
    """Parse a supported local file without modifying the source."""

    resolved = path.resolve(strict=True)
    if allowed_root is not None:
        root = allowed_root.resolve(strict=True)
        if not resolved.is_relative_to(root):
            raise ValueError("Document path is outside the allowed root")
    suffix = resolved.suffix.lower()
    if suffix == ".pdf":
        parsed = _parse_pdf(resolved)
    elif suffix == ".docx":
        parsed = _parse_docx(resolved)
    elif suffix == ".pptx":
        parsed = _parse_pptx(resolved)
    elif suffix == ".xlsx":
        parsed = _parse_xlsx(resolved)
    elif suffix in {".htm", ".html"}:
        parsed = _parse_html(resolved)
    elif suffix in {".md", ".markdown", ".txt"}:
        parsed = ParsedDocument(resolved.stem, resolved.stem, resolved.read_text(encoding="utf-8"))
    else:
        raise ValueError(f"Unsupported document type: {suffix or '<none>'}")
    if source_id is None:
        return parsed
    return ParsedDocument(source_id, parsed.title, parsed.text, parsed.page_spans)
