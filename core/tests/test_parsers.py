from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter

from loredock.ingestion import parse_path


def _write_pptx(path: Path) -> None:
    presentation = b"""<?xml version="1.0" encoding="UTF-8"?>
    <p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <p:sldIdLst><p:sldId id="256" r:id="rId2"/><p:sldId id="257" r:id="rId1"/></p:sldIdLst>
    </p:presentation>"""
    relationships = b"""<?xml version="1.0" encoding="UTF-8"?>
    <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rId1"
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
        Target="slides/slide1.xml"/>
      <Relationship Id="rId2"
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
        Target="slides/slide2.xml"/>
    </Relationships>"""
    slide_two = b"""<?xml version="1.0" encoding="UTF-8"?>
    <p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
      <p:cSld><p:spTree><p:sp><p:txBody>
        <a:p><a:r><a:t>First in presentation</a:t></a:r></a:p>
      </p:txBody></p:sp></p:spTree></p:cSld>
    </p:sld>"""
    slide_one = b"""<?xml version="1.0" encoding="UTF-8"?>
    <p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
      xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" show="0">
      <p:cSld><p:spTree><p:graphicFrame><a:graphic><a:graphicData><a:tbl>
        <a:tr><a:tc><a:txBody><a:p><a:r><a:t>Name</a:t></a:r></a:p></a:txBody></a:tc><a:tc><a:txBody><a:p><a:r><a:t>Value</a:t></a:r></a:p></a:txBody></a:tc></a:tr>
        <a:tr>
          <a:tc><a:txBody><a:p><a:r><a:t>LoreDock</a:t></a:r></a:p></a:txBody></a:tc>
          <a:tc><a:txBody><a:p><a:r><a:t>Local | first</a:t></a:r></a:p></a:txBody></a:tc>
        </a:tr>
      </a:tbl></a:graphicData></a:graphic></p:graphicFrame></p:spTree></p:cSld>
    </p:sld>"""
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("ppt/presentation.xml", presentation)
        archive.writestr("ppt/_rels/presentation.xml.rels", relationships)
        archive.writestr("ppt/slides/slide1.xml", slide_one)
        archive.writestr("ppt/slides/slide2.xml", slide_two)


def _write_xlsx(path: Path) -> None:
    workbook = b"""<?xml version="1.0" encoding="UTF-8"?>
    <workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <sheets>
        <sheet name="Summary" sheetId="1" r:id="rId2"/>
        <sheet name="Data" sheetId="2" state="hidden" r:id="rId1"/>
      </sheets>
    </workbook>"""
    relationships = b"""<?xml version="1.0" encoding="UTF-8"?>
    <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
      <Relationship Id="rId1"
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
        Target="worksheets/sheet1.xml"/>
      <Relationship Id="rId2"
        Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
        Target="worksheets/sheet2.xml"/>
    </Relationships>"""
    shared_strings = b"""<?xml version="1.0" encoding="UTF-8"?>
    <sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
      <si><t>Project</t></si><si><r><t>Lore</t></r><r><t>Dock</t></r></si>
    </sst>"""
    summary = b"""<?xml version="1.0" encoding="UTF-8"?>
    <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
      <sheetData><row r="1">
        <c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c>
        <c r="C1" t="inlineStr"><is><t>Local first</t></is></c>
      </row><row r="2">
        <c r="A2"><v>42</v></c><c r="B2" t="b"><v>1</v></c>
        <c r="C2"><f>SUM(A2,8)</f><v>50</v></c>
        <c r="D2"><f>UNAVAILABLE()</f></c>
      </row></sheetData>
    </worksheet>"""
    data = b"""<?xml version="1.0" encoding="UTF-8"?>
    <worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
      <sheetData><row r="3"><c r="B3" t="str"><v>Hidden value</v></c></row></sheetData>
    </worksheet>"""
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", relationships)
        archive.writestr("xl/sharedStrings.xml", shared_strings)
        archive.writestr("xl/worksheets/sheet1.xml", data)
        archive.writestr("xl/worksheets/sheet2.xml", summary)


def test_text_parser_rejects_path_outside_allowed_root(tmp_path: Path) -> None:
    allowed = tmp_path / "allowed"
    allowed.mkdir()
    outside = tmp_path / "outside.txt"
    outside.write_text("private", encoding="utf-8")

    with pytest.raises(ValueError, match="outside"):
        parse_path(outside, allowed_root=allowed)


def test_docx_parser_preserves_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("Second paragraph")
    document.save(str(path))

    parsed = parse_path(path)

    assert parsed.text == "第一段\n\nSecond paragraph"


def test_pdf_parser_records_page_spans(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as output:
        writer.write(output)

    parsed = parse_path(path)

    assert parsed.page_spans == ((1, 0, 0),)


def test_html_parser_extracts_structure_and_ignores_active_content(tmp_path: Path) -> None:
    path = tmp_path / "sample.html"
    path.write_text(
        """<!doctype html>
        <html><head><title>Lore &amp; Dock</title>
        <style>.secret { content: 'hidden'; }</style></head>
        <body><h1>知识中心</h1><iframe src="https://example.invalid/" />
        <p>本地优先 &amp; 可引用。</p>
        <ul><li>第一项</li><li>第二项</li></ul>
        <script>fetch('https://example.invalid/private')</script></body></html>""",
        encoding="utf-8",
    )

    parsed = parse_path(path, source_id="source-id")

    assert parsed.source_id == "source-id"
    assert parsed.title == "Lore & Dock"
    assert parsed.text == "知识中心\n\n本地优先 & 可引用。\n\n第一项\n第二项"
    assert "hidden" not in parsed.text
    assert "fetch" not in parsed.text


def test_html_parser_uses_filename_when_title_is_missing(tmp_path: Path) -> None:
    path = tmp_path / "notes.htm"
    path.write_text("<main>可搜索内容</main>", encoding="utf-8-sig")

    parsed = parse_path(path)

    assert parsed.title == "notes"
    assert parsed.text == "可搜索内容"


def test_pptx_parser_preserves_slide_order_tables_and_page_spans(tmp_path: Path) -> None:
    path = tmp_path / "deck.pptx"
    _write_pptx(path)

    parsed = parse_path(path, source_id="presentation-id")

    assert parsed.source_id == "presentation-id"
    assert parsed.title == "deck"
    assert parsed.text.startswith("# 幻灯片 1\n\nFirst in presentation")
    assert "# 幻灯片 2 (隐藏)" in parsed.text
    assert "| Name | Value |" in parsed.text
    assert "| LoreDock | Local \\| first |" in parsed.text
    assert tuple(page for page, _, _ in parsed.page_spans) == (1, 2)
    for page, start, end in parsed.page_spans:
        assert parsed.page_at(start) == page
        assert parsed.text[start:end].startswith(f"# 幻灯片 {page}")


def test_xlsx_parser_preserves_sheet_order_values_coordinates_and_page_spans(
    tmp_path: Path,
) -> None:
    path = tmp_path / "workbook.xlsx"
    _write_xlsx(path)

    parsed = parse_path(path, source_id="workbook-id")

    assert parsed.source_id == "workbook-id"
    assert parsed.title == "workbook"
    assert parsed.text.startswith("# 工作表 1: Summary")
    assert "A1: Project" in parsed.text
    assert "B1: LoreDock" in parsed.text
    assert "C1: Local first" in parsed.text
    assert "A2: 42" in parsed.text
    assert "B2: TRUE" in parsed.text
    assert "C2: 50" in parsed.text
    assert "D2: =UNAVAILABLE()" in parsed.text
    assert "# 工作表 2: Data (隐藏)" in parsed.text
    assert "B3: Hidden value" in parsed.text
    assert tuple(page for page, _, _ in parsed.page_spans) == (1, 2)
    for page, start, end in parsed.page_spans:
        assert parsed.page_at(start) == page
        assert parsed.text[start:end].startswith(f"# 工作表 {page}")
